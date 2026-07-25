from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "specification" / "noteflow-unified-spec-source.md"
OUTPUT = ROOT / "docs" / "specification" / "NoteFlow_统一架构与功能设计说明书_v2.0_整合版.docx"
ASSET_DIR = ROOT / "scripts" / ".noteflow_spec_assets"

CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "173A5E"
INK = "26313D"
MUTED = "66788A"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
BORDER = "D6DEE7"
WHITE = "FFFFFF"
FONT_ASCII = "Arial Unicode MS"
FONT_EAST_ASIA = "Arial Unicode MS"
FONT_MONO = "Menlo"


def set_run_font(run, *, size=None, bold=None, italic=None, color=None, name=FONT_ASCII, east_asia=FONT_EAST_ASIA):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa=TABLE_INDENT_DXA):
    if sum(widths_dxa) != CONTENT_DXA:
        raise ValueError(f"table widths must sum to {CONTENT_DXA}: {widths_dxa}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def paragraph_border(paragraph, *, side="left", color=BLUE, size=18, space=8):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    edge = p_bdr.find(qn(f"w:{side}"))
    if edge is None:
        edge = OxmlElement(f"w:{side}")
        p_bdr.append(edge)
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), str(space))
    edge.set(qn("w:color"), color)


def paragraph_shading(paragraph, fill: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    set_run_font(run, size=9, color=MUTED)


def _next_numbering_id(numbering, tag: str, attr: str) -> int:
    values = [int(node.get(qn(attr))) for node in numbering.findall(qn(tag))]
    return (max(values) + 1) if values else 1


def create_abstract_numbering(doc: Document, *, fmt: str, text: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_id = _next_numbering_id(numbering, "w:abstractNum", "w:abstractNumId")
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"{abstract_id:08X}")
    abstract.append(nsid)
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "271")
    p_pr.append(ind)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_ASCII)
    fonts.set(qn("w:hAnsi"), FONT_ASCII)
    fonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    r_pr.append(fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    return abstract_id


def create_numbering_instance(doc: Document, abstract_id: int) -> int:
    numbering = doc.part.numbering_part.element
    num_id = _next_numbering_id(numbering, "w:num", "w:numId")
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_id))
    num.append(abstract)
    numbering.append(num)
    return num_id


def set_paragraph_numbering(paragraph, num_id: int):
    p_pr = paragraph._p.get_or_add_pPr()
    old = p_pr.find(qn("w:numPr"))
    if old is not None:
        p_pr.remove(old)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_ASCII
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_ASCII)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_ASCII)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = FONT_ASCII
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_ASCII)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_ASCII)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = FONT_ASCII
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_ASCII)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_ASCII)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    caption = styles["Caption"]
    caption.font.name = FONT_ASCII
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)


def configure_section(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(table, [6500, 2860], indent_dxa=0)
    left = table.cell(0, 0).paragraphs[0]
    left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    left.paragraph_format.space_after = Pt(0)
    run = left.add_run("NOTEFLOW  /  统一架构与功能设计说明书")
    set_run_font(run, size=8.5, bold=True, color=MUTED)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_after = Pt(0)
    run = right.add_run("v2.0 整合版")
    set_run_font(run, size=8.5, color=MUTED)
    for cell in table.rows[0].cells:
        set_cell_margins(cell, 0, 0, 0, 0)
    table._tbl.tblPr.append(OxmlElement("w:tblBorders"))

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("NoteFlow  ·  ")
    set_run_font(run, size=9, color=MUTED)
    add_page_field(p)

    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""
    first_footer = section.first_page_footer
    first_footer.paragraphs[0].text = ""


def add_metadata_table(doc: Document):
    rows = [
        ("文档版本", "v2.0 整合版"),
        ("文档状态", "架构决策基线 / 统一实施依据"),
        ("核心范围", "产品能力、LangGraph、LlamaIndex RAG、Mem0、数据、API、目录、迁移、验收"),
        ("现状日期", "2026-07-17"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_geometry(table, [2400, 6960])
    set_table_borders(table, color=BORDER, size=5)
    for idx, (label, value) in enumerate(rows):
        c0, c1 = table.rows[idx].cells
        set_cell_shading(c0, LIGHT_BLUE)
        for cell, text, bold in ((c0, label, True), (c1, value, False)):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(text)
            set_run_font(run, size=10, bold=bold, color=NAVY if bold else INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_cover(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(46)
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run("NOTEFLOW  /  ARCHITECTURE BASELINE")
    set_run_font(run, size=10, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("NoteFlow 统一架构与\n功能设计说明书")
    set_run_font(run, size=28, bold=True, color=NAVY)
    paragraph_border(p, side="left", color=BLUE, size=28, space=12)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(30)
    run = p.add_run("整合 v2.0 RAG 升级方案、当前代码事实与后续模块化路线")
    set_run_font(run, size=14, color=MUTED)

    add_metadata_table(doc)

    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_DXA])
    set_table_borders(table, color="B9CAE0", size=6)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EEF5FB")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run("架构结论\n")
    set_run_font(run, size=11, bold=True, color=BLUE)
    run = p.add_run(
        "前端 UI 保持现状；LangGraph 负责全局流程，LlamaIndex 作为可调用 RAG 模块，Mem0 OSS 负责长期记忆。"
        "PostgreSQL + pgvector 继续作为统一数据底座，Redis 仅承担缓存、锁、限流和短暂运行状态。"
    )
    set_run_font(run, size=10.5, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("Prepared for NoteFlow product and engineering implementation")
    set_run_font(run, size=9.5, italic=True, color=MUTED)
    doc.add_page_break()


def extract_main_headings(lines: list[str]) -> list[str]:
    return [line[3:].strip() for line in lines if line.startswith("## ")]


def add_toc(doc: Document, headings: list[str]):
    p = doc.add_paragraph("目录", style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    lead = doc.add_paragraph("本文档按产品边界、目标架构、核心 AI 模块、工程实现和迁移验收组织。")
    lead.paragraph_format.space_after = Pt(10)
    # LibreOffice can continue numbering across distinct <w:num> instances that
    # share one abstract definition. Give the TOC its own abstract numbering so
    # every later ordered list can restart cleanly from 1.
    toc_abstract_id = create_abstract_numbering(doc, fmt="decimal", text="%1.")
    toc_num_id = create_numbering_instance(doc, toc_abstract_id)
    for heading in headings:
        p = doc.add_paragraph(style="List Number")
        set_paragraph_numbering(p, toc_num_id)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(re.sub(r"^\d+\.\s*", "", heading))
        set_run_font(run, size=9.3, color=INK)
    doc.add_page_break()


def font_path():
    candidates = [
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/Library/Fonts/Arial Unicode.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return candidate
    raise FileNotFoundError("No CJK font found")


def diagram_font(size):
    return ImageFont.truetype(font_path(), size=size)


def draw_arrow(draw, start, end, color="#557A99", width=5):
    draw.line([start, end], fill=color, width=width)
    x2, y2 = end
    if abs(end[0] - start[0]) > abs(end[1] - start[1]):
        sign = 1 if end[0] > start[0] else -1
        points = [(x2, y2), (x2 - sign * 16, y2 - 10), (x2 - sign * 16, y2 + 10)]
    else:
        sign = 1 if end[1] > start[1] else -1
        points = [(x2, y2), (x2 - 10, y2 - sign * 16), (x2 + 10, y2 - sign * 16)]
    draw.polygon(points, fill=color)


def draw_box(draw, xy, title, subtitle="", fill="#F4F7FA", outline="#9CB5C8", title_color="#173A5E"):
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = xy
    tf = diagram_font(30)
    sf = diagram_font(21)
    bbox = draw.textbbox((0, 0), title, font=tf)
    tx = (x1 + x2 - (bbox[2] - bbox[0])) / 2
    ty = y1 + 19
    draw.text((tx, ty), title, font=tf, fill=title_color)
    if subtitle:
        bbox = draw.multiline_textbbox((0, 0), subtitle, font=sf, spacing=6, align="center")
        sx = (x1 + x2 - (bbox[2] - bbox[0])) / 2
        sy = y1 + 66
        draw.multiline_text((sx, sy), subtitle, font=sf, fill="#526878", spacing=6, align="center")


def make_architecture_diagram(path: Path):
    img = Image.new("RGB", (1600, 980), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((70, 40), "NoteFlow 目标分层架构", font=diagram_font(42), fill="#173A5E")
    d.text((70, 100), "稳定前端契约，内部按编排、工具、能力与基础设施分层", font=diagram_font(24), fill="#66788A")
    layers = [
        ("前端与交互", "React / TypeScript / Tiptap / SSE", "#EEF5FB"),
        ("传输层", "FastAPI Routers / Auth / Pydantic / SSE Adapter", "#F5F8FA"),
        ("全局编排", "LangGraph Main Graph + Subgraphs + Interrupt", "#EAF2F8"),
        ("类型化工具", "Note / Draft / Edit / RAG / Memory Tools", "#F5F8FA"),
        ("AI 能力模块", "LlamaIndex RAG  |  Mem0 Memory  |  Model Providers", "#EEF5FB"),
        ("数据与运行", "PostgreSQL + pgvector  |  Redis  |  Workers  |  Object Storage", "#F5F8FA"),
    ]
    y = 160
    for idx, (title, subtitle, fill) in enumerate(layers):
        draw_box(d, (110, y, 1490, y + 105), title, subtitle, fill=fill)
        if idx < len(layers) - 1:
            draw_arrow(d, (800, y + 105), (800, y + 132))
        y += 132
    img.save(path)


def make_langgraph_diagram(path: Path):
    img = Image.new("RGB", (1600, 900), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((70, 40), "LangGraph 主图与受控子图", font=diagram_font(42), fill="#173A5E")
    boxes = {
        "ingress": (80, 160, 360, 270),
        "policy": (470, 160, 750, 270),
        "router": (860, 160, 1140, 270),
        "answer": (1240, 160, 1520, 270),
        "rag": (270, 430, 560, 560),
        "draft": (655, 430, 945, 560),
        "edit": (1040, 430, 1330, 560),
        "final": (655, 700, 945, 810),
    }
    draw_box(d, boxes["ingress"], "Ingress", "请求 / 会话 / Run")
    draw_box(d, boxes["policy"], "Policy", "UI 硬规则 / 安全")
    draw_box(d, boxes["router"], "Router", "Intent / Resume")
    draw_box(d, boxes["answer"], "Answer", "流式 LLM")
    draw_box(d, boxes["rag"], "RAG 子图", "LlamaIndex / 引用", fill="#EEF5FB")
    draw_box(d, boxes["draft"], "Draft 子图", "大纲 / Interrupt / Worker", fill="#FFF8E8", outline="#D7B867")
    draw_box(d, boxes["edit"], "Edit 子图", "Preview / Interrupt / Apply", fill="#FFF8E8", outline="#D7B867")
    draw_box(d, boxes["final"], "Finalize", "消息 / Trace / Memory Queue", fill="#EAF5EF", outline="#86B89D")
    draw_arrow(d, (360, 215), (470, 215))
    draw_arrow(d, (750, 215), (860, 215))
    draw_arrow(d, (1140, 215), (1240, 215))
    for x in (415, 800, 1185):
        draw_arrow(d, (1000, 270), (x, 430))
    for x in (415, 800, 1185, 1380):
        draw_arrow(d, (x, 560 if x != 1380 else 270), (800, 700))
    d.text((70, 845), "写操作只能通过 Draft/Edit 子图与确认节点；只读问答可走 RAG/Answer。", font=diagram_font(22), fill="#66788A")
    img.save(path)


def make_rag_diagram(path: Path):
    img = Image.new("RGB", (1600, 1000), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((70, 40), "企业级 RAG Pipeline v2.0", font=diagram_font(42), fill="#173A5E")
    draw_box(d, (90, 140, 470, 265), "Query Understanding", "Normalize / Rewrite / Expansion", fill="#EEF5FB")
    draw_box(d, (610, 140, 990, 265), "Permission Filter", "user_id / note_id / ACL", fill="#FFF4F1", outline="#D58B7B")
    draw_box(d, (1130, 140, 1510, 265), "Parallel Retrieval", "同一权限范围内并行", fill="#EEF5FB")
    draw_arrow(d, (470, 202), (610, 202))
    draw_arrow(d, (990, 202), (1130, 202))
    channels = [
        ("Title / Metadata", "Exact + pg_trgm", 90),
        ("BM25", "bm25s + 中文分词", 610),
        ("Vector", "DashScope + pgvector", 1130),
    ]
    for title, sub, x in channels:
        draw_box(d, (x, 385, x + 380, 520), title, sub, fill="#F5F8FA")
        draw_arrow(d, (1320, 265), (x + 190, 385))
        draw_arrow(d, (x + 190, 520), (800, 650))
    draw_box(d, (610, 650, 990, 770), "RRF Fusion", "去重 / 通道权重 / Top 20-50", fill="#EAF2F8")
    draw_box(d, (1110, 650, 1490, 770), "Reranker", "Cross Encoder / Top K", fill="#FFF8E8", outline="#D7B867")
    draw_arrow(d, (990, 710), (1110, 710))
    draw_box(d, (860, 845, 1240, 955), "Context Builder", "Token Budget / Citation / Strict Mode", fill="#EAF5EF", outline="#86B89D")
    draw_arrow(d, (1300, 770), (1050, 845))
    img.save(path)


def make_memory_diagram(path: Path):
    img = Image.new("RGB", (1600, 900), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((70, 40), "五层记忆与写入治理", font=diagram_font(42), fill="#173A5E")
    layers = [
        ("1  瞬时记忆", "当前输入 / 选区 / PageState", "AgentState，本轮"),
        ("2  短期记忆", "最近对话 / 会话摘要", "ChatSession，有预算"),
        ("3  工作记忆", "草稿 / 编辑 / 待确认", "Checkpoint，有 TTL"),
        ("4  情景记忆", "经历 / 项目事件", "Mem0 + 审计，时间衰减"),
        ("5  语义记忆", "稳定偏好 / 身份 / 长期目标", "Mem0 + 审计，冲突更新"),
    ]
    y = 145
    for i, (name, detail, storage) in enumerate(layers):
        draw_box(d, (80, y, 1020, y + 110), name, f"{detail}    |    {storage}", fill="#F5F8FA" if i < 3 else "#EEF5FB")
        y += 130
    draw_box(d, (1120, 190, 1520, 620), "Memory Policy", "显式性\n敏感/第三方\n临时性\n冲突/幂等\n用户开关\n当前输入优先", fill="#FFF8E8", outline="#D7B867")
    draw_arrow(d, (1020, 455), (1120, 405))
    draw_box(d, (1120, 680, 1520, 810), "异步写入", "Outbox -> Mem0 -> user_memories", fill="#EAF5EF", outline="#86B89D")
    draw_arrow(d, (1320, 620), (1320, 680))
    img.save(path)


def make_diagrams():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    files = {
        "architecture": ASSET_DIR / "architecture.png",
        "langgraph": ASSET_DIR / "langgraph.png",
        "rag": ASSET_DIR / "rag.png",
        "memory": ASSET_DIR / "memory.png",
    }
    make_architecture_diagram(files["architecture"])
    make_langgraph_diagram(files["langgraph"])
    make_rag_diagram(files["rag"])
    make_memory_diagram(files["memory"])
    return files


def add_diagram(doc: Document, path: Path, caption: str, index: int):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    inline = run.add_picture(str(path), width=Inches(6.3))
    doc_pr = inline._inline.docPr
    doc_pr.set("name", f"Figure {index}: {caption}")
    doc_pr.set("descr", caption)
    cp = doc.add_paragraph(f"图 {index}  {caption}", style="Caption")
    cp.paragraph_format.keep_with_next = False


def split_inline_code(paragraph, text: str, *, size=11, color=INK):
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size - 0.7, color=DARK_BLUE, name=FONT_MONO, east_asia=FONT_EAST_ASIA)
            paragraph_shading(paragraph, "F7F8FA")
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)


def add_callout(doc: Document, text: str):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_DXA])
    set_table_borders(table, color="B9CAE0", size=5)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EEF5FB")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.25
    split_inline_code(p, text, size=10.3)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 0.3


def choose_widths(columns: int) -> list[int]:
    if columns == 2:
        return [2700, 6660]
    if columns == 3:
        return [1900, 3250, 4210]
    if columns == 4:
        return [1550, 2450, 2500, 2860]
    return [CONTENT_DXA // columns] * (columns - 1) + [CONTENT_DXA - (CONTENT_DXA // columns) * (columns - 1)]


def add_markdown_table(doc: Document, rows: list[list[str]]):
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    set_table_geometry(table, choose_widths(columns))
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(rows):
        for c_idx in range(columns):
            cell = table.cell(r_idx, c_idx)
            text = row[c_idx].strip() if c_idx < len(row) else ""
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            split_inline_code(p, text, size=9.2 if columns >= 3 else 9.7, color=NAVY if r_idx == 0 else INK)
            if r_idx == 0:
                for run in p.runs:
                    run.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 0.5


def add_code_block(doc: Document, lines: list[str]):
    for idx, line in enumerate(lines or [""]):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.right_indent = Inches(0.05)
        p.paragraph_format.space_before = Pt(4 if idx == 0 else 0)
        p.paragraph_format.space_after = Pt(4 if idx == len(lines) - 1 else 0)
        p.paragraph_format.line_spacing = 1.05
        paragraph_shading(p, "F5F7FA")
        if idx == 0:
            paragraph_border(p, side="left", color="A9BED0", size=16, space=5)
        run = p.add_run(line if line else " ")
        set_run_font(run, size=8.4, color="2E3B48", name=FONT_MONO, east_asia=FONT_EAST_ASIA)


def should_page_break_before(text: str) -> bool:
    # Chapter 13 starts with a large table. Let Word/LibreOffice paginate that
    # heading naturally; a forced break there can create an extra blank page.
    return any(text.startswith(prefix) for prefix in ("4. ", "7. ", "9. ", "11. ", "16. ", "20. ", "24. "))


def render_markdown(
    doc: Document,
    lines: list[str],
    diagrams: dict[str, Path],
    bullet_abstract_id: int,
):
    i = 0
    figure_index = 1
    current_list_kind = None
    current_num_id = None
    while i < len(lines):
        raw = lines[i].rstrip("\n")
        line = raw.strip()
        if not line:
            current_list_kind = None
            current_num_id = None
            i += 1
            continue
        if line.startswith("# "):
            i += 1
            continue
        if line.startswith("## "):
            current_list_kind = None
            current_num_id = None
            text = line[3:].strip()
            if should_page_break_before(text) and len(doc.paragraphs) > 5:
                doc.add_page_break()
            p = doc.add_paragraph(text, style="Heading 1")
            p.paragraph_format.keep_with_next = True
            i += 1
            continue
        if line.startswith("### "):
            current_list_kind = None
            current_num_id = None
            doc.add_paragraph(line[4:].strip(), style="Heading 2")
            i += 1
            continue
        if line.startswith("#### "):
            current_list_kind = None
            current_num_id = None
            doc.add_paragraph(line[5:].strip(), style="Heading 3")
            i += 1
            continue
        if line.startswith("[DIAGRAM:"):
            current_list_kind = None
            current_num_id = None
            key = line[len("[DIAGRAM:") : -1]
            captions = {
                "architecture": "NoteFlow 目标总体架构与层次边界",
                "langgraph": "LangGraph 主图、只读路径与受控写入子图",
                "rag": "RAG v2.0：权限前置的混合检索、融合与重排",
                "memory": "五层记忆、Policy Engine 与异步写入",
            }
            add_diagram(doc, diagrams[key], captions[key], figure_index)
            figure_index += 1
            i += 1
            continue
        if line.startswith("```"):
            current_list_kind = None
            current_num_id = None
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i].rstrip("\n"))
                i += 1
            add_code_block(doc, block)
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[i + 1].strip()):
            current_list_kind = None
            current_num_id = None
            table_lines = [line]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = [[cell.strip() for cell in row.strip("|").split("|")] for row in table_lines]
            add_markdown_table(doc, rows)
            continue
        if line.startswith("> "):
            current_list_kind = None
            current_num_id = None
            add_callout(doc, line[2:].strip())
            i += 1
            continue
        if line.startswith("- "):
            if current_list_kind != "bullet":
                current_num_id = create_numbering_instance(doc, bullet_abstract_id)
                current_list_kind = "bullet"
            p = doc.add_paragraph(style="List Bullet")
            set_paragraph_numbering(p, current_num_id)
            split_inline_code(p, line[2:].strip())
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            if current_list_kind != "decimal":
                # A fresh abstract definition is intentional: it prevents office
                # renderers from continuing a previous list across sections.
                decimal_abstract_id = create_abstract_numbering(doc, fmt="decimal", text="%1.")
                current_num_id = create_numbering_instance(doc, decimal_abstract_id)
                current_list_kind = "decimal"
            text = re.sub(r"^\d+\.\s+", "", line)
            p = doc.add_paragraph(style="List Number")
            set_paragraph_numbering(p, current_num_id)
            split_inline_code(p, text)
            i += 1
            continue
        current_list_kind = None
        current_num_id = None
        p = doc.add_paragraph()
        split_inline_code(p, line)
        i += 1


def audit(doc: Document):
    for section in doc.sections:
        assert round(section.page_width.inches, 2) == 8.5
        assert round(section.page_height.inches, 2) == 11.0
        assert all(round(value.inches, 2) == 1.0 for value in (section.top_margin, section.right_margin, section.bottom_margin, section.left_margin))
    for table in doc.tables:
        grid = table._tbl.tblGrid
        widths = [int(col.get(qn("w:w"))) for col in grid]
        assert sum(widths) == CONTENT_DXA, widths


def main():
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    configure_styles(doc)
    configure_section(doc)
    doc.core_properties.title = "NoteFlow 统一架构与功能设计说明书"
    doc.core_properties.subject = "NoteFlow v2.0 整合架构、功能、实现与迁移设计"
    doc.core_properties.author = "NoteFlow Architecture"
    doc.core_properties.keywords = "NoteFlow, LangGraph, LlamaIndex, Mem0, RAG, pgvector"
    diagrams = make_diagrams()
    bullet_abstract_id = create_abstract_numbering(doc, fmt="bullet", text="•")
    add_cover(doc)
    add_toc(doc, extract_main_headings(lines))
    start = next(index for index, line in enumerate(lines) if line.startswith("## 0."))
    render_markdown(doc, lines[start:], diagrams, bullet_abstract_id)
    audit(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
